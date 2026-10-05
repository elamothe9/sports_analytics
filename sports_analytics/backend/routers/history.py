"""
Historical results endpoint — reads the props tracker spreadsheet live and
recomputes the same analysis the Excel Summary tab shows, so the web app can
display it without duplicating the data. All P&L is recomputed from the raw
odds + result cells (flat $1 stake), so it never depends on Excel having
recalculated its formulas.
"""
import os
from fastapi import APIRouter, HTTPException
import openpyxl

router = APIRouter()

# tk.xlsx lives in the repo root (two levels up from backend/routers/).
TRACKER_PATH = os.path.normpath(os.path.join(
    os.path.dirname(os.path.abspath(__file__)), "..", "..", "tk.xlsx"
))
SHEET = "Props Tracker"
DATA_START_ROW = 5

# display label -> (result column, odds column)  [1-based]
PROPS = [("HR", 9, 6), ("1H", 23, 22), ("2H", 10, 7), ("3H", 11, 8)]
PROP_LABELS = [p[0] for p in PROPS]
THRESHOLDS = [158, 160, 164, 167, 170, 172, 174, 176, 178, 180]


def _settled(v):
    return v in (0, 1)


def _as_odds(v):
    if v in (None, ""):
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _pnl(result, odds):
    """Flat $1 P&L for one prop. None = not settled."""
    if not _settled(result):
        return None
    if result == 0:
        return -1.0
    if odds is None:
        return None  # a win with no odds recorded — can't value it
    return odds / 100.0 if odds > 0 else 100.0 / abs(odds)


def _breakeven_pct(odds):
    if odds is None:
        return None
    return (100.0 / (odds + 100.0)) * 100 if odds > 0 \
        else (abs(odds) / (abs(odds) + 100.0)) * 100


@router.get("/history")
def history():
    if not os.path.exists(TRACKER_PATH):
        raise HTTPException(
            status_code=404,
            detail=f"Tracker spreadsheet not found at {TRACKER_PATH}",
        )
    try:
        wb = openpyxl.load_workbook(TRACKER_PATH, data_only=True)
    except PermissionError:
        raise HTTPException(
            status_code=423,
            detail="Tracker is locked — close tk.xlsx in Excel and retry.",
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Could not read tracker: {e}")

    if SHEET not in wb.sheetnames:
        raise HTTPException(status_code=500, detail=f"Sheet '{SHEET}' not found")
    ws = wb[SHEET]

    rows = []
    for r in range(DATA_START_ROW, ws.max_row + 1):
        player = ws.cell(r, 2).value
        date = ws.cell(r, 1).value
        if player is None or date is None:
            continue
        score = ws.cell(r, 4).value
        team = ws.cell(r, 3).value
        date_str = date.date().isoformat() if hasattr(date, "date") \
            else str(date)[:10]

        results, odds, pnl = {}, {}, {}
        day_pnl, day_has = 0.0, False
        for label, rc, oc in PROPS:
            rv = ws.cell(r, rc).value
            ov = _as_odds(ws.cell(r, oc).value)
            results[label] = rv if _settled(rv) else None
            odds[label] = ov
            p = _pnl(rv, ov)
            pnl[label] = round(p, 3) if p is not None else None
            if p is not None:
                day_pnl += p
                day_has = True

        rows.append({
            "date": date_str,
            "player": player,
            "team": team,
            "score": round(float(score), 1) if isinstance(score, (int, float)) else None,
            "results": results,
            "odds": odds,
            "pnl": pnl,
            "day_pnl": round(day_pnl, 2) if day_has else None,
        })

    # ---- per-prop summary ----
    summary = {}
    for label, rc, oc in PROPS:
        settled = [x for x in rows if x["results"][label] is not None]
        bets = len(settled)
        wins = sum(1 for x in settled if x["results"][label] == 1)
        pnl_sum = sum(x["pnl"][label] for x in settled if x["pnl"][label] is not None)
        odds_list = [x["odds"][label] for x in settled if x["odds"][label] is not None]
        avg_odds = sum(odds_list) / len(odds_list) if odds_list else None
        summary[label] = {
            "bets": bets,
            "wins": wins,
            "win_pct": round(wins / bets * 100, 1) if bets else None,
            "avg_odds": round(avg_odds) if avg_odds is not None else None,
            "breakeven_pct": round(_breakeven_pct(avg_odds), 1) if avg_odds is not None else None,
            "pnl": round(pnl_sum, 2),
            "roi": round(pnl_sum / bets * 100, 1) if bets else None,
        }

    # ---- profit / hit-rate by score threshold ----
    matrix = {label: [] for label in PROP_LABELS}
    for label, rc, oc in PROPS:
        for t in THRESHOLDS:
            sub = [x for x in rows
                   if x["results"][label] is not None
                   and isinstance(x["score"], (int, float)) and x["score"] >= t]
            bets = len(sub)
            wins = sum(1 for x in sub if x["results"][label] == 1)
            profit = sum(x["pnl"][label] for x in sub if x["pnl"][label] is not None)
            matrix[label].append({
                "cutoff": t,
                "bets": bets,
                "wins": wins,
                "win_pct": round(wins / bets * 100, 1) if bets else None,
                "roi": round(profit / bets * 100, 1) if bets else None,
                "profit": round(profit, 2),
            })

    dates = sorted({x["date"] for x in rows})
    log = sorted(rows, key=lambda x: (x["date"], x["player"]), reverse=True)

    return {
        "source": os.path.basename(TRACKER_PATH),
        "rows_logged": len(rows),
        "days_logged": len(dates),
        "date_range": {"start": dates[0], "end": dates[-1]} if dates else None,
        "props": PROP_LABELS,
        "thresholds": THRESHOLDS,
        "summary": summary,
        "matrix": matrix,
        "log": log,
    }
